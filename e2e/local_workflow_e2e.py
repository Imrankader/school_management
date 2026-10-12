#!/usr/bin/env python
"""
End-to-end workflow, data-reflection and role-permission test for the School Management backend.

SAFETY: this script CREATES, EDITS and DELETES records. It therefore refuses to run against anything
other than a local gateway (http://localhost:8080). Run it against the local stack
(C:\\Users\\<you>\\.tools\\school-local.ps1 start), never against the live site.

Auth: it signs short-lived JWTs for ADMIN / TEACHER / PARENT test users with the local dev secret
(school-local-jwt.secret), so no passwords are used. Test data is prefixed "E2E" and removed at the end
where the API offers a delete.

Usage:  python local_workflow_e2e.py [--base http://localhost:8080] [--keep]
"""
import base64
import datetime
import hashlib
import hmac
import io
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid

import openpyxl

BASE = "http://localhost:8080"
if "--base" in sys.argv:
    BASE = sys.argv[sys.argv.index("--base") + 1].rstrip("/")
if not (BASE.startswith("http://localhost") or BASE.startswith("http://127.0.0.1")):
    sys.exit(f"REFUSING to run: {BASE} is not a local address. This test writes data; use the local stack only.")

SECRET_FILE = os.environ.get("E2E_JWT_SECRET_FILE", r"C:\Users\yousuf.suhail_cloud-\.tools\school-local-jwt.secret")
SECRET = open(SECRET_FILE, "rb").read().strip()
RUN = datetime.datetime.now().strftime("%H%M%S")
TODAY = datetime.date.today()

results = []          # (group, name, ok, detail)
cleanup = []          # (description, callable)


# ------------------------------------------------------------------ helpers
def b64(x):
    return base64.urlsafe_b64encode(x).rstrip(b"=")


def make_token(role, user_id, email, **claims):
    now = int(time.time())
    head = b64(json.dumps({"alg": "HS256"}, separators=(",", ":")).encode())
    payload = {"sub": email, "role": role, "userId": user_id, "iat": now, "exp": now + 1800}
    payload.update({k: v for k, v in claims.items() if v is not None})
    body = b64(json.dumps(payload, separators=(",", ":")).encode())
    sig = b64(hmac.new(SECRET, head + b"." + body, hashlib.sha256).digest())
    return (head + b"." + body + b"." + sig).decode()


class Resp:
    def __init__(self, status, raw, headers):
        self.status, self.raw, self.headers = status, raw, headers
        try:
            self.json = json.loads(raw)
        except Exception:
            self.json = None

    @property
    def data(self):
        return self.json.get("data") if isinstance(self.json, dict) else None

    @property
    def ok(self):
        return 200 <= self.status < 300

    @property
    def denied(self):
        return self.status in (401, 403)

    def text(self, n=140):
        return self.raw.decode("utf-8", "replace")[:n].replace("\n", " ")


def call(method, path, token=None, body=None, files=None, params=None):
    url = BASE + path
    if params:
        url += ("&" if "?" in url else "?") + urllib.parse.urlencode(params)
    headers = {}
    data = None
    if token:
        headers["Authorization"] = "Bearer " + token
    if files:
        boundary = "----e2e" + uuid.uuid4().hex
        buf = io.BytesIO()
        for name, (fname, content, ctype) in files.items():
            buf.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"; filename=\"{fname}\"\r\n"
                      f"Content-Type: {ctype}\r\n\r\n".encode())
            buf.write(content)
            buf.write(b"\r\n")
        buf.write(f"--{boundary}--\r\n".encode())
        data = buf.getvalue()
        headers["Content-Type"] = "multipart/form-data; boundary=" + boundary
    elif body is not None:
        data = json.dumps(body).encode()
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(url, data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return Resp(r.status, r.read(), dict(r.headers))
    except urllib.error.HTTPError as e:
        return Resp(e.code, e.read(), dict(e.headers))


def check(group, name, ok, detail=""):
    results.append((group, name, bool(ok), str(detail)))
    return bool(ok)


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def class_summary_totals(token):
    r = call("GET", "/api/fees/class-summary", token)
    rows = r.data or []
    return (sum(float(x.get("totalAmount") or 0) for x in rows),
            sum(float(x.get("totalPaidAmount") or 0) for x in rows),
            sum(float(x.get("totalOutstandingAmount") or 0) for x in rows))


def student_counts(token):
    r = call("GET", "/api/students", token, params={"page": 1, "pageSize": 1})
    d = r.data or {}
    return d.get("totalStudents"), (d.get("pagination") or {}).get("total")


# ------------------------------------------------------------------ actors
ADMIN = make_token("ADMIN", 1, "e2e-admin@example.test")
TEACHER = make_token("TEACHER", 2, "e2e-teacher@example.test", assignedClass="LKG", assignedSection="A")

# ================================================================== 1. WORKFLOWS
G = "1 Student lifecycle"
before_total, before_pag = student_counts(ADMIN)
s1_adm, s1_name = f"E2E-{RUN}-1", f"E2E Alpha {RUN}"
r = call("POST", "/api/students", ADMIN, {
    "admissionNumber": s1_adm, "name": s1_name, "className": "LKG", "section": "A", "dateOfBirth": "2020-03-15",
    "gender": "Female", "fatherName": "E2E Father", "fatherMobileNumber": "9876500011", "contactNumber": "9876500011",
    "joiningDate": "2026-06-01", "isActive": True})
check(G, "admin creates a student (201)", r.status == 201, f"{r.status} {r.text(80)}")
S1 = (r.data or {}).get("id")
cleanup.append(("delete student S1", lambda: call("DELETE", f"/api/students/{S1}", ADMIN)))

r = call("GET", f"/api/students/{S1}", ADMIN)
check(G, "student detail shows saved date of birth", (r.data or {}).get("dateOfBirth") == "2020-03-15", (r.data or {}).get("dateOfBirth"))
check(G, "student detail shows saved class / section / mobile",
      all([(r.data or {}).get("className") == "LKG", (r.data or {}).get("section") == "A",
           (r.data or {}).get("fatherMobileNumber") == "9876500011"]))
r = call("GET", "/api/students", ADMIN, params={"search": s1_adm})
found = [s for s in ((r.data or {}).get("data") or []) if s.get("admissionNumber") == s1_adm]
check(G, "search by admission number finds the student", len(found) == 1, f"{len(found)} match")
t2, p2 = student_counts(ADMIN)
check(G, "dashboard total students +1", t2 == (before_total or 0) + 1 and p2 == (before_pag or 0) + 1, f"{before_total}->{t2}")
r = call("GET", "/api/students/class-summary", ADMIN)
lkg = next((x for x in (r.data or []) if (x.get("className") or "").upper() == "LKG"), {})
check(G, "class summary lists LKG with students", any(str(v).isdigit() and int(v) >= 1 for k, v in lkg.items() if "tudent" in k or "ount" in k), lkg)

# --- bulk import with the three DOB spellings customers type
G = "2 Bulk student import (DOB fix)"
r = call("GET", "/api/students/bulk-template", ADMIN, params={"className": "LKG"})
check(G, "template downloads as xlsx", r.ok and r.raw[:2] == b"PK", f"{r.status} {len(r.raw)} bytes")
wb = openpyxl.load_workbook(io.BytesIO(r.raw))
ws = wb.active
hdr = {str(c.value).strip(): i + 1 for i, c in enumerate(ws[1]) if c.value}
dob_col, join_col = hdr["DOB"], hdr["Joining Date"]
dv = [d for d in ws.data_validations.dataValidation if str(d.sqref).startswith(("F", "P"))]
check(G, "DOB / Joining cells do not carry a blocking date rule", dv and not any(d.showErrorMessage for d in dv),
      [(str(d.sqref), d.type, d.showErrorMessage) for d in dv])
rows_to_add = [
    (f"E2E-{RUN}-B1", f"E2E Bulk One {RUN}", "15-06-2021", "2021-06-15"),
    (f"E2E-{RUN}-B2", f"E2E Bulk Two {RUN}", "03/11/2021", "2021-11-03"),
    (f"E2E-{RUN}-B3", f"E2E Bulk Three {RUN}", datetime.date(2019, 12, 1), "2019-12-01"),
]
first_blank = next(i for i in range(2, ws.max_row + 1) if not ws.cell(i, hdr["Adm No"]).value)
for n, (adm, name, dob, _) in enumerate(rows_to_add):
    row = first_blank + n
    ws.cell(row, hdr["Adm No"]).value = adm
    ws.cell(row, hdr["Name"]).value = name
    ws.cell(row, hdr["Class"]).value = "LKG"
    ws.cell(row, hdr["Section"]).value = "A"
    ws.cell(row, dob_col).value = dob
    ws.cell(row, hdr["Father Name"]).value = "E2E Father"
    ws.cell(row, hdr["Father Mobile Number"]).value = f"98765{RUN[-5:]}"
    if isinstance(dob, datetime.date):
        ws.cell(row, dob_col).number_format = "dd-mm-yyyy"
buf = io.BytesIO()
wb.save(buf)
files = {"file": ("e2e_students.xlsx", buf.getvalue(), XLSX)}
r = call("POST", "/api/students/bulk-validate", ADMIN, files=files, params={"className": "LKG"})
check(G, "validation step accepts the sheet", (r.data or {}).get("success") is True, r.text(150))
r = call("POST", "/api/students/bulk-upload", ADMIN, files=files, params={"className": "LKG"})
d = r.data or {}
check(G, "upload adds exactly the 3 new students", d.get("insertedCount") == 3, {k: d.get(k) for k in ("insertedCount", "updatedCount", "errorCount", "message")})
bulk_ids = {}
for adm, name, _, want in rows_to_add:
    s = call("GET", "/api/students", ADMIN, params={"search": adm}).data
    s = next((x for x in (s or {}).get("data", []) if x["admissionNumber"] == adm), None)
    if s:
        bulk_ids[adm] = s["id"]
        cleanup.append((f"delete bulk student {adm}", (lambda i=s["id"]: call("DELETE", f"/api/students/{i}", ADMIN))))
    check(G, f"DOB saved for {adm.split('-')[-1]} ({want})", s and s.get("dateOfBirth") == want, s and s.get("dateOfBirth"))
S2 = bulk_ids.get(rows_to_add[0][0])
S2_name = rows_to_add[0][1]

G = "3 Billing, payments and totals"
tot0 = class_summary_totals(ADMIN)
r = call("POST", "/api/fees", ADMIN, {
    "studentId": S1, "studentName": s1_name, "admissionNumber": s1_adm, "className": "LKG",
    "termFees1": 4000, "termFees2": 4000, "termFees3": 2000, "busFees": 0, "examFees": 0,
    "totalAmount": 10000, "paidAmount": 0, "pendingAmount": 10000})
check(G, "admin creates a bill (201)", r.status == 201, f"{r.status} {r.text(90)}")
FEE = (r.data or {}).get("id")
r = call("POST", f"/api/fees/{FEE}/payments", ADMIN, {"amountPaid": 2500, "paymentDate": str(TODAY), "paymentMethod": "CASH", "note": "E2E part payment"})
check(G, "admin records a part payment (201)", r.status == 201, f"{r.status} {r.text(90)}")
r = call("GET", f"/api/fees/student/{S1}", ADMIN)
f = (r.data or [{}])[0]
check(G, "student bill: total 10000 / paid 2500 / due 7500 / PARTIAL",
      (f.get("totalAmount"), f.get("paidAmount"), f.get("pendingAmount"), f.get("status")) == (10000, 2500, 7500, "PARTIAL")
      or (float(f.get("totalAmount", 0)), float(f.get("paidAmount", 0)), float(f.get("pendingAmount", 0)), f.get("status")) == (10000.0, 2500.0, 7500.0, "PARTIAL"),
      {k: f.get(k) for k in ("totalAmount", "paidAmount", "pendingAmount", "status")})
tot1 = class_summary_totals(ADMIN)
check(G, "fee page / dashboard totals move by exactly +10000 billed, +2500 paid, +7500 due",
      [round(b - a) for a, b in zip(tot0, tot1)] == [10000, 2500, 7500], [round(b - a) for a, b in zip(tot0, tot1)])
r = call("GET", "/api/fees/class/LKG", ADMIN)
row = next((x for x in (r.data or []) if x.get("studentId") == S1), {})
check(G, "class billing table shows the student as 'partially paid'", row.get("status") == "partially paid", row.get("status"))
r = call("GET", "/api/fees/payments/recent", ADMIN, params={"limit": 20})
check(G, "recent payments lists the 2500 payment", any(p.get("studentId") == S1 and float(p.get("amountPaid")) == 2500 for p in (r.data or [])))
PARENT1 = make_token("PARENT", 3, "e2e-parent1@example.test", studentId=S1)
r = call("GET", f"/api/fees/student/{S1}", PARENT1)
check(G, "parent portal sees the same bill for their child", r.ok and r.data and float(r.data[0]["pendingAmount"]) == 7500, r.status)
r = call("POST", f"/api/fees/{FEE}/payments", ADMIN, {"amountPaid": 7500, "paymentDate": str(TODAY), "paymentMethod": "ONLINE", "note": "E2E settle"})
r = call("GET", f"/api/fees/student/{S1}", ADMIN)
check(G, "paying the balance settles the bill (PAID, due 0)", (r.data or [{}])[0].get("status") == "PAID" and float((r.data or [{}])[0].get("pendingAmount")) == 0, (r.data or [{}])[0].get("status"))

G = "4 Marks (new mark-sheet template)"
subs = {s["name"]: s["id"] for s in (call("GET", "/api/academic/subjects", ADMIN).data or [])}
tam, eng = subs.get("Tamil"), subs.get("English")
r = call("GET", "/api/academic/marks/bulk-template", ADMIN, params={"className": "LKG", "subjectIds": f"{tam},{eng}", "examName": "Quarterly"})
wb = openpyxl.load_workbook(io.BytesIO(r.raw))
ws = wb.active
head = [c.value for c in ws[1]]
check(G, "template columns: S.No, Mobile, Name, Class, subjects, Total, Average, Grade, absent flags",
      head[:4] == ["S.No", "Mobile", "Name", "Class"] and head[4:6] == ["Tamil", "English"] and head[6:9] == ["Total", "Average", "Grade"] and head[9:] == ["A Tamil", "A English"], head)
names = [ws.cell(i, 3).value for i in range(2, ws.max_row + 1)]
check(G, "template is pre-filled with the LKG students (incl. the new one)", s1_name in names, f"{len(names)} students")
row = 2 + names.index(s1_name)
ws.cell(row, 5).value = 80          # Tamil
ws.cell(row, 11).value = "A"        # English absent
buf = io.BytesIO()
wb.save(buf)
r = call("POST", "/api/academic/marks/bulk-upload", ADMIN, params={"examName": "Quarterly"},
         files={"file": ("e2e_marks.xlsx", buf.getvalue(), XLSX)})
d = r.data or {}
check(G, "upload: 1 mark + 1 absent saved, no errors", (d.get("successful"), d.get("failed")) == (2, 0), {k: d.get(k) for k in ("successful", "updated", "failed", "message")})
mk = {m["subjectName"]: m for m in (call("GET", "/api/academic/marks", ADMIN, params={"className": "LKG", "examName": "Quarterly"}).data or []) if m.get("admissionNumber") == s1_adm}
check(G, "marks view: Tamil 80/100 grade A", mk.get("Tamil", {}).get("marksObtained") == 80 and mk["Tamil"].get("grade") == "A", mk.get("Tamil", {}).get("grade"))
check(G, "marks view: English absent -> grade AB, remark Absent", mk.get("English", {}).get("grade") == "AB" and mk["English"].get("remarks") == "Absent")
r = call("GET", f"/api/academic/marks/student/{S1}", PARENT1)
check(G, "parent portal shows both marks", r.ok and len(r.data or []) == 2, f"{len(r.data or [])} records")
r = call("PUT", f"/api/academic/marks/{mk['Tamil']['id']}", TEACHER, {**mk["Tamil"], "marksObtained": 95})
after = call("GET", f"/api/academic/marks/student/{S1}", PARENT1).data or []
t = next((m for m in after if m["subjectName"] == "Tamil"), {})
check(G, "teacher edits Tamil to 95 -> parent portal reflects it with new grade", t.get("marksObtained") == 95 and t.get("grade") == "A+", {k: t.get(k) for k in ("marksObtained", "grade")})

G = "5 Leave"
r = call("POST", "/api/attendance/leave", PARENT1, {"leaveDate": str(TODAY), "startDate": str(TODAY), "endDate": str(TODAY + datetime.timedelta(days=1)), "reason": f"E2E fever {RUN}"})
check(G, "parent submits a leave (201/200)", r.ok, f"{r.status} {r.text(80)}")
adm_leaves = call("GET", "/api/attendance/leave", ADMIN).data or []
mine = [x for x in adm_leaves if x.get("reason") == f"E2E fever {RUN}"]
check(G, "admin leave page lists it with the student's name", len(mine) == 1 and mine[0].get("studentName") == s1_name, mine and mine[0].get("studentName"))
check(G, "...and class / section", mine and (mine[0].get("className"), mine[0].get("section")) == ("LKG", "A"))
tl = call("GET", "/api/attendance/leave/teacher", TEACHER).data or []
check(G, "teacher of LKG-A sees it", any(x.get("reason") == f"E2E fever {RUN}" for x in tl), f"{len(tl)} records")
PARENT2 = make_token("PARENT", 4, "e2e-parent2@example.test", studentId=S2)
my2 = call("GET", "/api/attendance/leave/my", PARENT2).data or []
check(G, "another parent does NOT see it", not any(x.get("reason") == f"E2E fever {RUN}" for x in my2))
my1 = call("GET", "/api/attendance/leave/my", PARENT1).data or []
check(G, "the submitting parent sees it", any(x.get("reason") == f"E2E fever {RUN}" for x in my1))
TEACHER_OTHER = make_token("TEACHER", 5, "e2e-teacher2@example.test", assignedClass="UKG", assignedSection="A")
tl2 = call("GET", "/api/attendance/leave/teacher", TEACHER_OTHER).data or []
check(G, "teacher of a different class does NOT see it", not any(x.get("reason") == f"E2E fever {RUN}" for x in tl2))

G = "6 Homework and holidays"
r = call("POST", "/api/academic/homework", TEACHER, {"title": f"E2E worksheet {RUN}", "description": "page 1", "className": "LKG", "section": "A", "subjectId": tam, "dueDate": str(TODAY + datetime.timedelta(days=3))})
check(G, "teacher assigns homework", r.ok, f"{r.status} {r.text(80)}")
HW = (r.data or {}).get("id")
hw = call("GET", "/api/academic/homework/class/LKG", ADMIN).data or []
check(G, "homework appears for the class (admin view)", any(h["id"] == HW for h in hw))
r = call("PUT", f"/api/academic/homework/{HW}", TEACHER, {**(next((h for h in hw if h["id"] == HW), {})), "title": f"E2E worksheet edited {RUN}"})
hw = call("GET", "/api/academic/homework", PARENT1).data or []
check(G, "edit is visible to a parent", any(h["id"] == HW and h["title"].endswith(f"edited {RUN}") for h in hw))
call("DELETE", f"/api/academic/homework/{HW}", TEACHER)
check(G, "deleted homework disappears", not any(h["id"] == HW for h in (call("GET", "/api/academic/homework", ADMIN).data or [])))
r = call("POST", "/api/attendance/holidays", ADMIN, {"name": f"E2E Holiday {RUN}", "date": str(TODAY + datetime.timedelta(days=20)), "description": "test", "holidayType": "PUBLIC"})
HOL = (r.data or {}).get("id")
check(G, "admin adds a holiday", r.ok, r.status)
check(G, "teacher sees the holiday", any(h["id"] == HOL for h in (call("GET", "/api/attendance/holidays", TEACHER).data or [])))
call("DELETE", f"/api/attendance/holidays/{HOL}", ADMIN)
check(G, "deleted holiday disappears", not any(h["id"] == HOL for h in (call("GET", "/api/attendance/holidays", ADMIN).data or [])))

G = "7 Notifications"
msgs = {"STUDENTS": f"E2E to parents {RUN}", "TEACHERS": f"E2E to teachers {RUN}", "BOTH": f"E2E to everyone {RUN}"}
for aud, m in msgs.items():
    call("POST", "/api/notifications", ADMIN, {"date": str(TODAY), "audience": aud, "message": m})
pm = {n["message"] for n in (call("GET", "/api/notifications/my", PARENT1).data or [])}
tm = {n["message"] for n in (call("GET", "/api/notifications/my", TEACHER).data or [])}
hm = {n["message"] for n in (call("GET", "/api/notifications/history", ADMIN).data or [])}
check(G, "parent gets STUDENTS + BOTH only", msgs["STUDENTS"] in pm and msgs["BOTH"] in pm and msgs["TEACHERS"] not in pm)
check(G, "teacher gets TEACHERS + BOTH only", msgs["TEACHERS"] in tm and msgs["BOTH"] in tm and msgs["STUDENTS"] not in tm)
check(G, "admin history shows all three", set(msgs.values()) <= hm)

# ================================================================== 2. PERMISSION MATRIX
# "FAIL" in this section = a role can do something it should not be able to do (security gap).
G = "8 Who can do what"
PARENT_UNLINKED = make_token("PARENT", 9, "e2e-selfregistered@example.test")   # what public /register creates
FULL = {"className": "LKG", "section": "A", "gender": "Male", "dateOfBirth": "2020-01-01", "fatherName": "E2E Father",
        "fatherMobileNumber": "9876500099", "contactNumber": "9876500099", "joiningDate": "2026-06-01", "isActive": True}
tw_resp = call("POST", "/api/students", ADMIN, {**FULL, "admissionNumber": f"E2E-{RUN}-X", "name": f"E2E Throwaway {RUN}"})
TW = (tw_resp.data or {}).get("id")
check(G, "test setup: throwaway student created by admin", TW is not None, f"{tw_resp.status} {tw_resp.text(100)}")
if TW is None:
    sys.exit("cannot continue the permission matrix without the throwaway student")
cleanup.append(("delete throwaway student", lambda: call("DELETE", f"/api/students/{TW}", ADMIN)))
exam_id = (call("GET", "/api/academic/exams", ADMIN).data or [{}])[0].get("id")


def deny(who, label, resp, created=None):
    rule_rejection = resp.status == 400 and "Only PARENT" in resp.text(400)
    ok = resp.denied or rule_rejection
    why = "allowed -- NOT BLOCKED" if resp.ok else ("blocked by business rule" if rule_rejection else "")
    check(G, f"{who}: {label}", ok, f"HTTP {resp.status} {why}".strip())
    if resp.ok and created:
        cleanup.append((f"remove what {who} created ({label})", created(resp)))


P = PARENT1
deny("PARENT (own child S1)", "cannot read ANOTHER student's profile", call("GET", f"/api/students/{S2}", P))
deny("PARENT (own child S1)", "cannot list ALL students", call("GET", "/api/students", P, params={"pageSize": 1}))
deny("PARENT (own child S1)", "cannot export all students", call("GET", "/api/students/export", P))
deny("PARENT (own child S1)", "cannot read ANOTHER student's bill", call("GET", f"/api/fees/student/{S2}", P))
deny("PARENT (own child S1)", "cannot read ANOTHER student's payments", call("GET", f"/api/fees/student/{S2}/payments", P))
deny("PARENT (own child S1)", "cannot read school-wide fee totals", call("GET", "/api/fees/class-summary", P))
deny("PARENT (own child S1)", "cannot read ANOTHER student's marks", call("GET", f"/api/academic/marks/student/{S2}", P))
deny("PARENT (own child S1)", "cannot read all marks of an exam", call("GET", f"/api/academic/marks/exam/{exam_id}", P))
deny("PARENT (own child S1)", "cannot read the school-wide leave list", call("GET", "/api/attendance/leave", P))
deny("PARENT (own child S1)", "cannot create a student", call("POST", "/api/students", P, {**FULL, "admissionNumber": f"E2E-{RUN}-P", "name": f"E2E ParentMade {RUN}"}),
     lambda resp: (lambda i=(resp.data or {}).get("id"): call("DELETE", f"/api/students/{i}", ADMIN)))
deny("PARENT (own child S1)", "cannot edit a student", call("PUT", f"/api/students/{TW}", P, {**FULL, "admissionNumber": f"E2E-{RUN}-X", "name": f"E2E Hacked {RUN}"}))
deny("PARENT (own child S1)", "cannot change a student's status", call("PATCH", f"/api/students/{TW}/status", P, {"status": "INACTIVE"}))
deny("PARENT (own child S1)", "cannot delete a student", call("DELETE", f"/api/students/{TW}", P))
deny("PARENT (own child S1)", "cannot create a bill", call("POST", "/api/fees", P, {"studentId": S2, "studentName": S2_name, "className": "LKG", "totalAmount": 1, "paidAmount": 0, "pendingAmount": 1}))
deny("PARENT (own child S1)", "cannot edit a bill", call("PUT", f"/api/fees/{FEE}", P, {"studentId": S1, "totalAmount": 1, "paidAmount": 1, "pendingAmount": 0}))
deny("PARENT (own child S1)", "cannot record a payment", call("POST", f"/api/fees/{FEE}/payments", P, {"amountPaid": 1, "paymentDate": str(TODAY), "paymentMethod": "CASH"}))
deny("PARENT (own child S1)", "cannot add a holiday", call("POST", "/api/attendance/holidays", P, {"name": f"E2E parent holiday {RUN}", "date": str(TODAY + datetime.timedelta(days=30)), "holidayType": "PUBLIC"}),
     lambda resp: (lambda i=(resp.data or {}).get("id"): call("DELETE", f"/api/attendance/holidays/{i}", ADMIN)))
deny("PARENT (own child S1)", "cannot add a leave reason", call("POST", "/api/attendance/leave-reasons", P, {"englishReason": f"E2E reason {RUN}", "tamilMeaning": "x"}))
deny("PARENT (own child S1)", "cannot enter marks", call("POST", "/api/academic/marks", P, {"studentId": S1, "subjectName": "Tamil", "examName": "Quarterly", "marksObtained": 100, "maxMarks": 100}))
deny("PARENT (own child S1)", "cannot post homework", call("POST", "/api/academic/homework", P, {"title": "x", "className": "LKG", "section": "A", "dueDate": str(TODAY)}))
deny("PARENT (own child S1)", "cannot broadcast a notification", call("POST", "/api/notifications", P, {"date": str(TODAY), "audience": "BOTH", "message": "x"}))
deny("PARENT (own child S1)", "cannot read notification history", call("GET", "/api/notifications/history", P))
deny("PARENT (own child S1)", "cannot create an exam", call("POST", "/api/academic/exams", P, {"name": "E2E Exam", "totalMarks": 100}))
deny("PARENT (own child S1)", "cannot create a TEACHER/ADMIN user", call("POST", "/api/auth/register", P, {"name": "E2E x", "email": f"e2e{RUN}@example.test", "password": "E2e-test-pass-1", "role": "TEACHER"}))
deny("PARENT (own child S1)", "cannot list parent accounts", call("GET", "/api/auth/parents", P))
U = PARENT_UNLINKED
deny("PARENT (self-registered, no child)", "cannot read any student's marks", call("GET", f"/api/academic/marks/student/{S1}", U))
deny("PARENT (self-registered, no child)", "cannot read any student's profile", call("GET", f"/api/students/{S1}", U))
deny("PARENT (self-registered, no child)", "cannot read any student's bill", call("GET", f"/api/fees/student/{S1}", U))
deny("PARENT (self-registered, no child)", "cannot list ALL students", call("GET", "/api/students", U, params={"pageSize": 1}))
T = TEACHER
deny("TEACHER", "cannot create a bill", call("POST", "/api/fees", T, {"studentId": S2, "studentName": S2_name, "className": "LKG", "totalAmount": 1, "paidAmount": 0, "pendingAmount": 1}))
deny("TEACHER", "cannot record a payment", call("POST", f"/api/fees/{FEE}/payments", T, {"amountPaid": 1, "paymentDate": str(TODAY), "paymentMethod": "CASH"}))
deny("TEACHER", "cannot read school-wide fee totals", call("GET", "/api/fees/class-summary", T))
deny("TEACHER", "cannot delete a student", call("DELETE", f"/api/students/{TW}", T))
deny("TEACHER", "cannot broadcast a notification", call("POST", "/api/notifications", T, {"date": str(TODAY), "audience": "BOTH", "message": "x"}))
deny("TEACHER", "cannot delete marks", call("DELETE", f"/api/academic/marks/{mk['English']['id']}", T))
deny("TEACHER", "cannot create a TEACHER/ADMIN user", call("POST", "/api/auth/register", T, {"name": "E2E x", "email": f"e2e2{RUN}@example.test", "password": "E2e-test-pass-1", "role": "ADMIN"}))
deny("ANYONE (no token)", "cannot call the API", call("GET", "/api/students"))

# ================================================================== cleanup + report
KEEP = "--keep" in sys.argv          # leave the E2E records in place so the pages can be inspected in a browser
if KEEP:
    print("\n--keep: E2E records left in the local database. Test ids:",
          json.dumps({"run": RUN, "S1": S1, "S1_name": s1_name, "S2": S2, "S2_name": S2_name, "fee": FEE}))
    cleanup.clear()
for desc, fn in reversed(cleanup):
    try:
        fn()
    except Exception as e:  # keep going
        print("cleanup problem:", desc, e)

w = max(len(r[1]) for r in results) + 2
cur, fails = None, []
for g, n, ok, d in results:
    if g != cur:
        print(f"\n[{g}]")
        cur = g
    print(f"  {'PASS' if ok else 'FAIL'}  {n:<{w}} {d if (not ok or len(str(d)) < 60) else ''}")
    if not ok:
        fails.append((g, n, d))
print(f"\n{len(results) - len(fails)}/{len(results)} checks passed, {len(fails)} failed")
sys.exit(1 if fails else 0)
